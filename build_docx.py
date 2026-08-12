#!/usr/bin/env python3
"""
Render cv_content.py to an editable, ATS-friendly Word document.

    python3 build_docx.py
"""

import re
from html import unescape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

import cv_content as C

OUT = "AZODO_EMEKA_CV_Rotating_Equipment_Engineer_QatarEnergyLNG.docx"

NAVY = RGBColor(0x12, 0x30, 0x4F)
ACCENT = RGBColor(0x1F, 0x5C, 0x8B)
GREY = RGBColor(0x3F, 0x4A, 0x54)
INK = RGBColor(0x1A, 0x1A, 0x1A)
FONT = "Calibri"
BODY_PT = 9.0
CELL_PT = 8.5

TAG_RE = re.compile(r"(<b>|</b>)")


def clean(text):
    """Strip the tiny markup vocabulary shared with the PDF renderer."""
    return unescape(text.replace("&nbsp;", " "))


def add_runs(par, text, size=BODY_PT, color=INK, bold_all=False, italic=False):
    """Write text into a paragraph, honouring <b>...</b> spans."""
    bold = bold_all
    for piece in TAG_RE.split(text):
        if piece == "<b>":
            bold = True
        elif piece == "</b>":
            bold = bold_all
        elif piece:
            run = par.add_run(clean(piece))
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.bold = bold
            run.italic = italic
    return par


def spacing(par, before=0, after=2, line=1.0):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    return par


# WordprocessingML enforces a fixed child order inside pPr/tblPr; appending
# elements at the end produces a file Word and LibreOffice refuse to open.
PPR_AFTER_SHD = (
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
TBLPR_AFTER_BORDERS = (
    "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
    "w:tblDescription",
)


def shade(par, hex_color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hex_color)
    par._p.get_or_add_pPr().insert_element_before(el, *PPR_AFTER_SHD)


def section_heading(doc, title):
    par = doc.add_paragraph()
    spacing(par, before=5, after=2)
    shade(par, "12304F")
    add_runs(par, title.upper(), size=10, color=RGBColor(0xFF, 0xFF, 0xFF),
             bold_all=True)


def body_par(doc, text):
    par = doc.add_paragraph()
    spacing(par, after=2.5)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(par, text)


def bullet_par(doc, text):
    par = doc.add_paragraph(style="List Bullet")
    spacing(par, after=1.5)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.left_indent = Cm(0.5)
    add_runs(par, text)


def sub_par(doc, text):
    par = doc.add_paragraph()
    spacing(par, before=5, after=1)
    add_runs(par, text, size=BODY_PT, color=ACCENT, bold_all=True)


def role_par(doc, title, dates, widths=(Cm(12.7), Cm(4.7))):
    """Job title left, dates right — a borderless two-column row, so a long
    title wraps within its own column instead of running into the dates."""
    spacer = doc.add_paragraph()
    spacing(spacer, after=0)
    spacer.paragraph_format.line_spacing = Pt(2)

    table = doc.add_table(rows=1, cols=2)
    cells = table.rows[0].cells
    left = cells[0].paragraphs[0]
    spacing(left, after=0)
    add_runs(left, title, size=10, color=NAVY, bold_all=True)
    right = cells[1].paragraphs[0]
    spacing(right, after=0)
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_runs(right, dates, size=8.5, color=GREY, bold_all=True)
    _borders(table, hide=True)
    _fixed_layout(table, widths)


def org_par(doc, text):
    par = doc.add_paragraph()
    spacing(par, after=2)
    add_runs(par, text, size=8.5, color=GREY, italic=True)


def _fixed_layout(table, widths):
    """Word ignores cell widths unless the table layout is fixed."""
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.insert_element_before(
        layout, "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
    margins = OxmlElement("w:tblCellMar")
    for side, val in (("top", 0), ("left", 0), ("bottom", 0), ("right", 80)):
        m = OxmlElement("w:%s" % side)
        m.set(qn("w:w"), str(val))
        m.set(qn("w:type"), "dxa")
        margins.append(m)
    table._tbl.tblPr.insert_element_before(
        margins, "w:tblLook", "w:tblCaption", "w:tblDescription")
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    for col, width in zip(table.columns, widths):
        col.width = width


def kv_table(doc, rows):
    widths = (Cm(4.0), Cm(13.4))
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for label, value in rows:
        cells = table.add_row().cells
        p0 = cells[0].paragraphs[0]
        spacing(p0, after=1)
        add_runs(p0, label, size=CELL_PT, color=NAVY, bold_all=True)
        p1 = cells[1].paragraphs[0]
        spacing(p1, after=1)
        add_runs(p1, value, size=CELL_PT)
    _borders(table)
    _fixed_layout(table, widths)


def grid_table(doc, items):
    widths = (Cm(5.8),) * 3
    rows = [items[i:i + 3] for i in range(0, len(items), 3)]
    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in rows:
        cells = table.add_row().cells
        for i in range(3):
            par = cells[i].paragraphs[0]
            spacing(par, after=1)
            if i < len(row):
                add_runs(par, "• " + row[i], size=CELL_PT)
    _borders(table, hide=True)
    _fixed_layout(table, widths)


def _borders(table, hide=False):
    """Light horizontal rules only (or none), instead of Word's default box."""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:%s" % edge)
        visible = (not hide) and edge == "insideH"
        el.set(qn("w:val"), "single" if visible else "none")
        el.set(qn("w:sz"), "4" if visible else "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "B8C4CE")
        borders.append(el)
    table._tbl.tblPr.insert_element_before(borders, *TBLPR_AFTER_BORDERS)


def main():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY_PT)
    style.paragraph_format.space_after = Pt(2)

    sec = doc.sections[0]
    sec.top_margin = Cm(1.15)
    sec.bottom_margin = Cm(1.15)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)

    # ---- header -------------------------------------------------------
    par = doc.add_paragraph()
    spacing(par, after=1)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(par, C.NAME, size=19, color=NAVY, bold_all=True)

    par = doc.add_paragraph()
    spacing(par, after=2)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(par, C.TAGLINE, size=10.5, color=ACCENT, bold_all=True)

    for line in C.CONTACT:
        par = doc.add_paragraph()
        spacing(par, after=0)
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(par, line, size=8.5, color=GREY)

    # ---- sections -----------------------------------------------------
    for title, blocks in C.SECTIONS:
        section_heading(doc, title)
        for block in blocks:
            kind = block[0]
            if kind == "para":
                body_par(doc, block[1])
            elif kind == "bullet":
                bullet_par(doc, block[1])
            elif kind == "sub":
                sub_par(doc, block[1])
            elif kind == "role":
                _, rtitle, org, dates = block
                if title == "Education":
                    role_par(doc, rtitle, dates, widths=(Cm(8.6), Cm(8.8)))
                else:
                    role_par(doc, rtitle, dates)
                if org:
                    org_par(doc, org)
            elif kind == "kv":
                kv_table(doc, block[1])
            elif kind == "grid":
                grid_table(doc, block[1])
            else:
                raise ValueError("unknown block type: %r" % kind)

    # ---- footer -------------------------------------------------------
    fpar = sec.footer.paragraphs[0]
    fpar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(fpar, C.FOOTER_NAME, size=8, color=GREY)

    doc.core_properties.author = "Chukwuemeka Osmund Azodo"
    doc.core_properties.title = ("Chukwuemeka Osmund Azodo - CV - "
                                 "Rotating Equipment Engineer")
    doc.core_properties.subject = ("Application: Rotating Equipment Engineer, "
                                   "QatarEnergy LNG (Ras Laffan)")
    doc.save(OUT)
    print("Wrote %s" % OUT)


if __name__ == "__main__":
    main()
